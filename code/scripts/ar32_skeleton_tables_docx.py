"""ar32_skeleton_tables_docx.py — 排出回复导师用的新表（导师原规格）。

输入（只读）：本 exploration 的 output/tables/ 下
  ar26_mentor_spec.csv（企业层）；ar29_iv_dyad_mentor.csv（工具变量、双边）；
  ar31_skeleton_fill.csv（KP LM、第一阶段、机制两步、异质性交互、五项联合）；ar31_desc.csv（描述统计）
格式：master_supporting_docs/journal_benchmarks_20260924/表格格式裁决.md；排版函数见 ar22_fanwen_lib.py
输出：本 exploration 的 output/tables/回复导师_新表_第八轮_20261007.docx
运行：python3 explorations/advisor_revision_20261005/scripts/ar32_skeleton_tables_docx.py（项目根目录）
"""
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import ar22_fanwen_lib as L  # noqa: E402

T = HERE.parent / "output/tables"
OUT = T / "回复导师_新表_第八轮_20261007.docx"
L.VARS.update({"Inv", "Cashf", "Lservice", "Lerner", "WW", "ASY", "Seg", "RDres", "Srisk", "Resil", "Investnn",
               "Investnc", "Investnns", "LcomRDp", "LindRDs", "RDexp", "DuNum3", "DuRate3", "DuNum5", "DuRate5",
               "LMDA", "Pri_Number", "Pri_Hold"})
m = pd.read_csv(T / "ar26_mentor_spec.csv", dtype={"col": str})
n = pd.read_csv(T / "ar29_iv_dyad_mentor.csv", dtype={"col": str})
f = pd.read_csv(T / "ar31_skeleton_fill.csv", dtype={"col": str})
ds = pd.read_csv(T / "ar31_desc.csv")


def stars(p):
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


def get(d, tab, col, var):
    r = d[(d.tab == tab) & (d.col == str(col)) & (d["var"] == var)]
    assert len(r) == 1, (tab, col, var)
    return r.iloc[0]


def pair(r, k=1, dec=4):
    return f"{float(r.b) * k:.{dec}f}{stars(float(r.p))}", f"({float(r.se) * k:.{dec}f})"


def rows(label, rs, k=1):
    ps = [pair(r, k) for r in rs]
    return [[label] + [x[0] for x in ps], [""] + [x[1] for x in ps]]


def sparse(label, cells, k=1):
    """cells：每列为一行结果或 None（该列不含此变量）。"""
    ps = [pair(r, k) if r is not None else ("", "") for r in cells]
    return [[label] + [x[0] for x in ps], [""] + [x[1] for x in ps]]


def stats(rs, r2=True):
    out = [["样本量"] + [str(int(float(r.N))) for r in rs]]
    if r2:
        out.append(["调整后R²"] + [f"{float(r.r2a):.4f}" for r in rs])
    return out


def para(doc, text):
    p = doc.add_paragraph()
    L.exact_spacing(p, 18, before=12, after=0)
    L.set_fonts(p.add_run(text), 9, east="宋体")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = 7560310, 10692130
sec.top_margin = sec.bottom_margin = 914400
sec.left_margin = sec.right_margin = 1090295
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
yes = L.yes
FE2 = lambda k: [["企业固定效应"] + yes(k), ["行业×年份固定效应"] + yes(k)]
FE1 = lambda k: [["企业固定效应"] + yes(k), ["年份固定效应"] + yes(k)]
FED = lambda k: [["企业×年份固定效应"] + yes(k), ["企业×目的省固定效应"] + yes(k), ["目的省×年份固定效应"] + yes(k)]
f4 = lambda x: L.fix_minus(f"{x:.4f}")

# 补充包说明：表2、初稿8表3其余列与Heckman沿用初稿8（D-130 第4点）
p0 = doc.add_paragraph()
L.exact_spacing(p0, 18, before=0, after=6)
L.set_fonts(p0.add_run("本文件为补充表，与初稿8合读。表2基准回归、初稿8表3的其余稳健性列和Heckman两阶段结果沿用初稿8。"), 9, east="宋体")

# 表1 新增变量描述统计
L.caption(doc, "表1 新增变量描述性统计")
lab = {"Investnns": "同省异市新增", "entry": "当年新设（双边）", "Seg0_z": "目的地市场分割（标准化）", "east": "东部",
       "mkt0": "市场化程度", "rdres0": "研发资源", "tfp0": "期初生产率", "lqw": "同行业区位熵"}
body = [[lab.get(r["var"], r["var"]), str(int(r.N)), f4(r["mean"]), f4(r.sd), f4(r["min"]), f4(r.p50), f4(r["max"])]
        for _, r in ds.iterrows()]
L.build(doc, [["变量", "样本量", "均值", "标准差", "最小值", "中位数", "最大值"]], body, label_w=2400, keep=False)
L.note(doc, "注：各变量取其所在回归的估计样本。初稿8表1已有变量不再重复。Investnns为新增同省异市子公司数量加1取对数；功能类型新增均为该类新增"
       "跨省子公司数量加1取对数；双边变量取企业、目的省、年份三维面板2015年以后的样本，研发型新增（双边）为当年是否在该省新设研发型子公司。")

para(doc, "表2 基准回归结果沿用初稿8表2，此处不重复。")

# 表3 新增跨省投资与子公司存续
L.caption(doc, "表3 新增跨省投资与子公司存续")
a = [get(m, "T2", c, "Patient") for c in (1, 2, 3)]
s = [get(m, "T9", c, "xlag") for c in (1, 2, 3, 4)]
body = sparse("Patient", a + [None] * 4) + sparse("L.Patient", [None] * 3 + s)
body += [["控制变量"] + yes(7)] + FE2(7) + stats(a + s)
L.build(doc, [L.nums_head(7), ["", "新增", "新增", "新增", "存续", "存续", "存续", "存续"],
              ["", "Investnn", "Investnc", "Investnns", "DuNum3", "DuRate3", "DuNum5", "DuRate5"]], body, label_w=1900)
L.note(doc, "注：第（1）—（3）列对应初稿8表32第（1）—（3）列。第（4）—（7）列以基期跨省子公司存量为基础，取第三、五年仍在年报中披露（持续披露）的子公司数量与比例，"
       "要求企业在第三、五年仍在样本中，解释变量按初稿8写法取上一期；与原表样本量8127相符的候选口径见附表5。" + L.STD)

# 表4 稳健性检验
L.caption(doc, "表4 稳健性检验")
rs = [get(m, "T1", c, "Patient") for c in (1, 2)]
body = rows("Patient", rs) + [["控制变量"] + yes(2)] + FE2(2) + stats(rs)
L.build(doc, [L.nums_head(2), ["", "全样本", "剔除直辖市总部企业"], ["", "Investp", "Investp"]], body, label_w=3000)
L.note(doc, "注：第（1）列为基准回归，对应初稿8表2第（3）列；第（2）列剔除总部位于北京、上海、天津、重庆的企业。初稿8表3的替换被解释变量、"
       "被解释变量滞后一期、替换解释变量、剔除特大超大城市、控制国家区域战略政策、PSM、安慰剂检验各列与本表合并，合并后不超过8列。" + L.STD)

# 表5 内生性检验
L.caption(doc, "表5 内生性检验")
iv = [get(n, "T7", 1, "Patient_lag1"), get(n, "T7", 2, "lp"), get(n, "T7", 3, "lp")]
kp = [get(f, "I", c, "kp") for c in (1, 2, 3)]
body = rows("L.Patient", iv) + [["控制变量"] + yes(3), ["企业固定效应"] + yes(3), ["年份固定效应"] + yes(3),
                                ["样本量"] + [str(int(float(r.N))) for r in iv],
                                ["Kleibergen–Paap rk LM统计量"] + [f"{float(r.stat):.4f}" for r in kp],
                                ["Kleibergen–Paap rk Wald F"] + [f"{float(r.kpf):.4f}" for r in iv]]
L.build(doc, [L.nums_head(3), ["", "同群工具", "PRI签署基金数", "PRI签署基金持股"], ["", "Investp", "Investp", "Investp"]],
        body, label_w=3000)
L.note(doc, "注：报告第二阶段，内生变量为上一期Patient，第一阶段见附表1。第（1）列以初稿8所用滞后同群留一工具Hold(t−1)与SizeG(t−1)为工具；"
       "第（2）（3）列以上一期签署联合国负责任投资原则（PRI）的基金数、持股比例为工具，控制变量取上一期值。Kleibergen–Paap rk LM统计量的p值均小于0.01。"
       "Heckman两阶段结果沿用初稿8表3。" + L.STD)

# 表6 机制检验
L.caption(doc, "表6 机制检验")
rs = [get(m, "T3", c, "Patient") for c in (1, 2, 3, 5, 6, 7)]
body = rows("Patient", rs) + [["控制变量"] + yes(6)] + FE2(6) + stats(rs)
L.build(doc, [L.nums_head(6), ["", "WW", "ASY", "Srisk", "LcomRDp", "LindRDs", "RDexp"]], body, label_w=1800)
L.note(doc, "注：WW为融资约束指数，ASY为信息不对称，Srisk为供应链风险，LcomRDp为母公司创新，LindRDs为子公司创新，RDexp为集团创新地理分散度。"
       "LcomRDp取母公司专利总数（独立与联合专利获得量之和）加1的对数，是对原表指标的候选替代，不等同于独立申请口径；LindRDs取子公司独立专利总和的对数。LMDA数据暂缺，补齐后放在第（1）列之前。Resil见附表4，第二步见附表2。" + L.STD)

# 表7 异质性
L.caption(doc, "表7 异质性分析")
rs = [get(m, "T4", c, "Patient") for c in range(1, 9)]
pr = lambda g: f"{float(m[(m.tab == 'T4') & (m.col == g)].p.iloc[0]):.4f}"
body = rows("Patient", rs) + [["组间系数差异p值", "", pr("Seg"), "", pr("Dist"), "", pr("Resi"), "", pr("RD")],
                              ["控制变量"] + yes(8)] + FE1(8) + stats(rs)
head = [L.nums_head(8), ["", "市场分割", "市场分割", "地理距离", "地理距离", "供应链韧性", "供应链韧性", "研发投入强度", "研发投入强度"],
        ["", "高", "低", "远", "近", "高", "低", "高", "低"]]
L.build(doc, head, body, label_w=1640)
L.note(doc, "注：被解释变量为Investp。市场分割、供应链韧性按中位数分组，地理距离按球面距离分组，研发投入强度按当年中位数分组。"
       "组间系数差异p值按初稿8的检验写法（两组共用企业与年份固定效应），统一采用企业层面聚类；该检验中控制变量的系数也按组变化。交互项估计见附表3。" + L.STD)

# 表8 新增子公司的去向
L.caption(doc, "表8 新增子公司的去向")
rs = [get(n, "T6", 1, "PxSeg"), get(n, "T6", 2, "PxSeg")] + [get(n, "T8", c, f"Px_{c}") for c in range(1, 6)]
body = rows("L.Patient×目的省属性", rs, k=1000) + FED(7) + stats(rs)
L.build(doc, [L.nums_head(7), ["", "市场分割", "市场分割", "东部", "市场化\n程度", "研发资源", "期初\n生产率", "同行业\n区位熵"],
              ["", "全样本", "剔除\n直辖市", "全样本", "全样本", "全样本", "全样本", "全样本"]], body, label_w=2000)
L.note(doc, "注：企业、目的省、年份三维面板，被解释变量为当年是否在该省新设跨省子公司，每列放入一个目的省属性，第（2）列剔除直辖市目的地。市场分割为目的地市场分割；"
       "市场化程度、研发资源取2010—2013年均值，期初生产率为2014年该省上市公司TFP均值，同行业区位熵按2010—2013年该省本行业上市公司占比计算；"
       "除东部外均已标准化。系数与标准误均乘以1000。" + L.STD)

# 表9 去向属性的联合估计与研发型新增
L.caption(doc, "表9 去向属性的联合估计与研发型新增")
J = {a: get(f, "J", 1, f"c.L_Patient#c.{a}") for a in ("east", "z_mkt0", "z_rdres0", "z_tfp0", "z_lqw")}
rd = get(n, "T6", 3, "PxRD")
body = []
for a, labx in (("east", "东部"), ("z_mkt0", "市场化程度"), ("z_rdres0", "研发资源"), ("z_tfp0", "期初生产率"), ("z_lqw", "同行业区位熵")):
    body += sparse(f"L.Patient×{labx}", [J[a], rd if a == "z_rdres0" else None], k=1000)
body += FED(2) + stats([J["east"], rd])
L.build(doc, [L.nums_head(2), ["", "当年新增", "研发型新增"]], body, label_w=3000)
L.note(doc, "注：第（1）列把表8第（3）—（7）列的五项目的省属性放入同一回归；第（2）列被解释变量为当年是否在该省新设研发型子公司，研发型按第三方标签"
       "剔除房地产子公司后认定。期初生产率换用剔除样本少于10家省份的算法后的结果见附表6。系数与标准误均乘以1000。" + L.STD)

# 表10 新增跨省子公司的功能类型
L.caption(doc, "表10 新增跨省子公司的功能类型")
rs = [get(m, "T5", c, "Patient") for c in (1, 2, 3, 4, 6, 5)]
body = rows("Patient", rs) + [["控制变量"] + yes(6)] + FE2(6) + stats(rs)
L.build(doc, [L.nums_head(6), ["", "进入新省份", "同功能复制", "功能互补", "研发型", "研发互补", "东部企业到中西部"]], body, label_w=1950)
L.note(doc, "注：被解释变量为该类新增跨省子公司数量加1取对数。进入新省份指企业上一年在该省没有子公司；同功能复制指上一年在该省已有功能相同的子公司；"
       "功能互补指上一年在该省已有子公司但功能都不同；研发型按第三方标签剔除房地产子公司后认定；研发互补指研发型新子公司进入上一年已有生产或销售子公司、"
       "但没有研发子公司的省份；第（6）列样本为总部在东部的企业，被解释变量为其在中西部新设的子公司。各类子公司基数不同，系数不宜直接比较大小。" + L.STD)

# 附表1 第一阶段
L.caption(doc, "附表1 内生性检验第一阶段")
g = lambda c, v: get(f, "F1", c, v)
body = sparse("Hold(t−1)", [g(1, "IV_Hold2_lag1"), None, None]) + sparse("SizeG(t−1)", [g(1, "IV_SizeGroup2_lag1"), None, None])
body += sparse("L.Pri_Number", [None, g(2, "l_Pri_Number"), None]) + sparse("L.Pri_Hold", [None, None, g(3, "l_Pri_Hold")])
body += [["控制变量"] + yes(3), ["企业固定效应"] + yes(3), ["年份固定效应"] + yes(3)]
body += stats([g(1, "IV_Hold2_lag1"), g(2, "l_Pri_Number"), g(3, "l_Pri_Hold")])
L.build(doc, [L.nums_head(3), ["", "同群工具", "PRI签署基金数", "PRI签署基金持股"], ["", "L.Patient", "L.Patient", "L.Patient"]],
        body, label_w=3000)
L.note(doc, "注：对应表5各列的第一阶段，样本与表5相同。" + L.STD)

# 附表2 机制两步
L.caption(doc, "附表2 机制检验第二步")
mv = ["WW指数", "ASY", "SCDRisk2_100倍", "M_comA", "l子公司独立专利总和", "集团创新地理分散度"]
pp = [get(f, "A2", c, "Patient") for c in range(1, 7)]
mm = [get(f, "A2", c, v) for c, v in zip(range(1, 7), mv)]
body = rows("Patient", pp) + rows("机制变量", mm) + [["控制变量"] + yes(6)] + FE2(6) + stats(pp)
L.build(doc, [L.nums_head(6), ["", "WW", "ASY", "Srisk", "LcomRDp", "LindRDs", "RDexp"],
              ["", "Investp", "Investp", "Investp", "Investp", "Investp", "Investp"]], body, label_w=1800)
L.note(doc, "注：被解释变量为Investp，各列同时放入Patient和表头所列机制变量。" + L.STD)

# 附表3 异质性交互项
L.caption(doc, "附表3 异质性分析交互项估计")
cs = range(1, 5)
body = rows("Patient", [get(f, "A3", c, "Patient") for c in cs]) + rows("Patient×高组", [get(f, "A3", c, "PxG") for c in cs])
body += rows("高组", [get(f, "A3", c, "G") for c in cs]) + [["控制变量"] + yes(4)] + FE1(4) + stats([get(f, "A3", c, "PxG") for c in cs])
L.build(doc, [L.nums_head(4), ["", "市场分割高", "地理距离远", "供应链韧性高", "研发投入强度高"], ["", "Investp", "Investp", "Investp", "Investp"]],
        body, label_w=2200)
L.note(doc, "注：高组为虚拟变量，分组定义同表7。本表只让Patient按组变化，表7的组间差异检验同时让控制变量按组变化，两者显著性不完全相同。" + L.STD)

# 附表4 供应链韧性机制
L.caption(doc, "附表4 供应链韧性机制")
r = [get(m, "T3", 4, "Patient")]
body = rows("Patient", r) + [["控制变量"] + yes(1)] + FE2(1) + stats(r)
L.build(doc, [L.nums_head(1), ["", "Resil"]], body, label_w=3000, full_w=5000)
L.note(doc, "注：规格同表6。" + L.STD)

# 附表5 存续的原表样本口径
L.caption(doc, "附表5 第三年持续披露数量的候选样本口径")
r = [get(m, "T9", 5, "xlag")]
body = rows("L.Patient", r) + [["控制变量"] + yes(1)] + FE2(1) + stats(r)
L.build(doc, [L.nums_head(1), ["", "DuNum3"]], body, label_w=3000, full_w=5000)
L.note(doc, "注：第三年已不在样本中的企业记为0，其余同表3第（4）列。该口径的样本量与初稿8表32的8127相同，但仅样本量相符，尚不能确认是原表的构造方法。" + L.STD)

# 附表6 期初生产率换算法
L.caption(doc, "附表6 去向属性联合估计的期初生产率换算法")
J2 = [get(f, "J", 2, f"c.L_Patient#c.{a}") for a in ("east", "z_mkt0", "z_rdres0", "z_tfp0_n10", "z_lqw")]
body = []
for r, labx in zip(J2, ("东部", "市场化程度", "研发资源", "期初生产率", "同行业区位熵")):
    body += rows(f"L.Patient×{labx}", [r], k=1000)
body += FED(1) + stats([J2[0]])
L.build(doc, [L.nums_head(1), ["", "当年新增"]], body, label_w=3000, full_w=5000)
L.note(doc, "注：同表9第（1）列，期初生产率改为剔除上市公司样本少于10家的省份后计算。系数与标准误均乘以1000。" + L.STD)

# 页脚页码（D-130：方便面谈定位）
from docx.oxml import OxmlElement  # noqa: E402
fp = sec.footer.paragraphs[0]
fp.alignment = 1
for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    r = fp.add_run()
    L.set_fonts(r, 9)
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._element.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._element.append(e)
doc.save(OUT)
print(OUT)

# 只供本机导出 PDF 的副本：macOS 的 LibreOffice 不经 fontconfig，找不到“宋体/黑体”时会用 Arial Unicode 代替；
# 副本把东亚字体换成本机的宋体-简、黑体-简，Word 交付件仍写宋体、黑体（范文规定）。
import zipfile  # noqa: E402
PDFSRC = T / "回复导师_新表_第八轮_20261007_pdf用.docx"
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(PDFSRC, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename.endswith(".xml"):
            data = data.decode("utf-8").replace('"宋体"', '"Songti SC"').replace('"黑体"', '"Heiti SC"').encode("utf-8")
        zo.writestr(it, data)
print(PDFSRC)
