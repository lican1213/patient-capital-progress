"""ar88_t8v2_docx.py — 第四十七轮：新表8（导师清单版）Word，范文三线表格式，A4 纵向（由 ar85_t8_docx.py 改写）。
输入（只读）：本 exploration 的 output/tables/ar88_t8v2.csv
输出：deliverables/advisor_package_20261007/表8新版_20261012/表8_新版.docx；本 exploration 的 output/tables/表8_新版_pdf用.docx（渲染用）
运行：python3 explorations/advisor_revision_20261005/scripts/ar88_t8v2_docx.py（项目根目录）
"""
import sys
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import ar22_fanwen_lib as L  # noqa: E402

ROOT = HERE.parents[2]
T = HERE.parent / "output/tables"
OUT = ROOT / "deliverables/advisor_package_20261007/表8新版_20261012/表8_新版.docx"
L.VARS.update({"Patient"})
d = pd.read_csv(T / "ar88_t8v2.csv")


def stars(p):
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


def cell(col, var):
    r = d[(d.col == col) & (d["var"] == var)]
    if len(r) == 0:
        return "", ""
    assert len(r) == 1
    r = r.iloc[0]
    return f"{r.b * 1000:.3f}{stars(r.p)}", f"({r.se * 1000:.3f})"


ROWS = [("Patient×市场分割", "PxSeg"), ("市场分割", "seg"), ("Patient×东部", "PxEast"), ("Patient×中部", "PxCent"),
        ("Patient×市场化程度", "PxMkt"), ("Patient×研发资源", "PxRd"), ("Patient×GDP之比", "PxGdp"), ("Patient×地理距离", "PxDist")]
body = []
for lab, v in ROWS:
    cs = [cell(c, v) for c in range(1, 9)]
    body += [[lab] + [x[0] for x in cs], [""] + [x[1] for x in cs]]
first = d.drop_duplicates("col").set_index("col").sort_index()
body += [["企业×年份固定效应"] + L.yes(8), ["企业×目的省固定效应"] + L.yes(8), ["目的省×年份固定效应"] + L.yes(8),
         ["样本量"] + [str(int(first.loc[c, "N"])) for c in range(1, 9)],
         ["调整后R²"] + [f"{first.loc[c, 'r2a']:.4f}" for c in range(1, 9)]]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = 7560310, 10692130
sec.top_margin = sec.bottom_margin = 914400
sec.left_margin = sec.right_margin = 1090295
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

L.caption(doc, "表8 新增子公司的去向")
L.build(doc, [L.nums_head(8),
              ["", "市场分割", "市场分割", "东中西", "市场化\n程度", "研发\n资源", "GDP\n之比", "地理\n距离", "同时\n放入"],
              ["", "全样本", "剔除\n直辖市", "全样本", "全样本", "全样本", "全样本", "全样本", "全样本"]], body, label_w=2100, full_w=9200)
L.note(doc, "注：企业、目的省、年份三维面板，被解释变量为当年是否在该省新设跨省子公司；样本为2015年起，各企业首个样本年不计入；第（2）列剔除直辖市目的地。"
       "市场分割为母公司所在城市（取企业首个样本年）与目的省内各城市之间的城市对市场分割指数（基于价格指数），省内等权平均后取前三年（t−3至t−1）均值，"
       "再换算为当年全部观测中的百分位（0至1）；东部、中部按国家统计局东中西三分，以西部为对照；市场化程度、研发资源取当年值，研发资源为该省上市公司发明专利申请量合计加1取对数，二者在估计样本内标准化；"
       "GDP之比为目的省当年地区生产总值与母公司所在省当年地区生产总值之比取对数；地理距离为母公司所在城市到目的省内各城市球面距离的等权平均取对数，不加权。"
       "市场分割随企业、目的省和年份同时变化，其本身一并放入；其余特征本身被固定效应吸收，只放交互项。系数与标准误均乘以1000。" + L.STD)

fp = sec.footer.paragraphs[0]
fp.alignment = 1
for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    r = fp.add_run()
    L.set_fonts(r, 9)
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._element.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._element.append(e)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
pdfsrc = T / "表8_新版_pdf用.docx"
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(pdfsrc, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename.endswith(".xml"):
            data = data.decode("utf-8").replace('"宋体"', '"Songti SC"').replace('"黑体"', '"Heiti SC"').encode("utf-8")
        zo.writestr(it, data)
print(OUT); print(pdfsrc)
